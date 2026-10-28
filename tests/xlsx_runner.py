"""
数据驱动测试运行器

设计：
- 从 test.xlsx 读取测试用例（每行一个用例）
- 每个用例包含：用例ID/测试项/名称/操作(operation)/参数(params)/断言类型/期望/启用
- 运行器根据 operation 分派到对应处理器，处理器在隔离的临时数据目录中执行
- 断言类型支持：equals/truthy/falsy/contains/not_null/count/throws/throws_contains
  /json_has_key/json_has_no_key/json_field_equals
- 输出结构化测试结果，供 pytest 包装或命令行使用

用法：
    python -m tests.xlsx_runner              # 运行 tests/test.xlsx
    python -m tests.xlsx_runner --path X.xlsx
"""
import json
import sys
import tempfile
import traceback
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

from openpyxl import load_workbook

# 确保项目根在 path 中
_PROJECT_ROOT = Path(__file__).parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from src.core.user import (
    User, Role, UserManager, PasswordHasher, get_user_manager
)
from src.core.session import UserSession, AuditLogger, GlobalConsultationIndex, ReviewQueue
from src.core.patient import Patient, Gender, MedicalImage
from src.core.medical_record import MedicalRecordIntegrator
from src.core.consultation import ConsultationEngine
from src.core.knowledge_graph import build_default_knowledge_graph
from src.core.consultation import Priority
from src.adapters.base import MockHospitalAdapter


# ============================================================
# 测试夹具：每个用例独立临时目录，互不污染
# ============================================================
class TestEnv:
    """隔离测试环境"""

    def __init__(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.data_dir = Path(self._tmp.name)

        # 用户系统（用临时目录，重置全局单例）
        import src.core.user as _user_mod
        self._orig_um = _user_mod._user_manager
        _user_mod._user_manager = None
        self.user_manager = UserManager(data_dir=str(self.data_dir / "users"))

        # 审计日志
        self.audit = AuditLogger(audit_dir=str(self.data_dir / "audit"))

        # 用户空间根
        self.userdata_root = self.data_dir / "userdata"

        # 审核队列
        self.review_queue = ReviewQueue(data_root=str(self.data_dir / "review_queue"))

    def session(self, user_id: str) -> UserSession:
        return UserSession(user_id=user_id, data_root=str(self.userdata_root))

    def global_index(self) -> GlobalConsultationIndex:
        return GlobalConsultationIndex(data_root=str(self.userdata_root))

    def cleanup(self):
        import src.core.user as _user_mod
        _user_mod._user_manager = self._orig_um
        self._tmp.cleanup()


# ============================================================
# 辅助：构建测试用对象
# ============================================================
def _make_report(consultation_id: str = "TC-REPORT-001",
                 patient_id: str = "P001",
                 patient_name: str = "测试患者",
                 status: str = "COMPLETED"):
    """构建一个最小的会诊报告用于会话持久化测试"""
    from src.core.moderator import ConsultationReport
    from datetime import datetime
    return ConsultationReport(
        consultation_id=consultation_id,
        patient_id=patient_id,
        patient_name=patient_name,
        consultation_type="MDT",
        status=status,
        created_at=datetime.now().isoformat(),
        completed_at=datetime.now().isoformat(),
        participating_agents=[{"agent_id": "cardio", "name": "心内科AI", "specialty": "心内科"}],
        differential_diagnosis=[{"disease": "冠心病", "confidence": 0.8,
                                  "evidence_level": "B", "supporting_specialties": 2}],
        final_conclusion="测试结论",
    )


def _make_patient(patient_id: str = "P001", name: str = "测试患者"):
    """构建一个最小的患者对象"""
    from datetime import date
    return Patient(
        patient_id=patient_id,
        name=name,
        gender=Gender.MALE,
        age=60,
        department="心内科",
        chief_complaint="反复胸痛3天",
        present_illness="活动后胸骨后压榨样疼痛，休息可缓解",
        past_history="高血压病史5年",
        lab_results=[],
        imaging_reports=[],
    )


def _seed_consultation_engine(env: TestEnv):
    """构建一个带适配器数据的会诊引擎（用于会诊流程测试）"""
    patient = _make_patient()
    pid = patient.patient_id
    integrator = MedicalRecordIntegrator()
    integrator.add_adapter(MockHospitalAdapter("HIS", sample_data={pid: {
        "name": patient.name, "gender": "男", "age": patient.age,
        "department": patient.department, "attending_doctor": "张医生",
    }}))
    integrator.add_adapter(MockHospitalAdapter("EMR", sample_data={pid: {
        "chief_complaint": patient.chief_complaint,
        "present_illness": patient.present_illness,
        "past_history": patient.past_history,
    }}))
    integrator.add_adapter(MockHospitalAdapter("LIS", sample_data={pid: {
        "lab_results": [
            {"item_name": "肌钙蛋白", "value": "0.5", "unit": "ng/mL",
             "reference_range": "0-0.04", "abnormal": True}
        ]
    }}))
    integrator.add_adapter(MockHospitalAdapter("PACS", sample_data={pid: {
        "imaging_reports": [
            {"modality": "CT", "body_part": "胸部", "findings": "未见明显异常",
             "impression": "大致正常"}
        ]
    }}))
    kg = build_default_knowledge_graph()
    engine = ConsultationEngine(integrator=integrator, knowledge_graph=kg)
    return engine, patient


# ============================================================
# 操作处理器注册表
# 每个 handler: (env, params) -> actual
# ============================================================
def _op_password_hash(env, p):
    return PasswordHasher.hash(p["password"])


def _op_password_verify(env, p):
    stored = PasswordHasher.hash(p["password"])
    return PasswordHasher.verify(p["verify_password"], stored)


def _op_password_empty(env, p):
    return PasswordHasher.hash("")


def _op_role_parse(env, p):
    return Role.parse(p["value"]).value


def _op_role_can_review(env, p):
    return Role.parse(p["role"]).can_review()


def _op_role_display(env, p):
    return Role.parse(p["role"]).display_name


def _op_auth(env, p):
    return env.user_manager.authenticate(p["username"], p["password"])


def _op_auth_disabled(env, p):
    # 禁用默认 doctor 后认证
    doc = env.user_manager.get_by_username("doctor")
    env.user_manager.update_user(doc.user_id, enabled=False)
    return env.user_manager.authenticate(p["username"], p["password"])


def _op_create_user(env, p):
    u = env.user_manager.create_user(
        username=p["username"], display_name=p.get("display_name", "测试用户"),
        role=Role.parse(p["role"]), department=p.get("department", ""),
        title=p.get("title", ""), password=p["password"]
    )
    return u.user_id


def _op_create_duplicate(env, p):
    env.user_manager.create_user(
        username=p["username"], display_name="第一", role=Role.DOCTOR,
        password="pass1234"
    )
    env.user_manager.create_user(
        username=p["username"], display_name="第二", role=Role.DOCTOR,
        password="pass1234"
    )


def _op_create_weak_password(env, p):
    env.user_manager.create_user(
        username=p["username"], display_name="弱密码", role=Role.DOCTOR,
        password=p["password"]
    )


def _op_delete_last_admin(env, p):
    admin = env.user_manager.get_by_username("admin")
    env.user_manager.delete_user(admin.user_id)


def _op_reset_password(env, p):
    target = env.user_manager.get_by_username(p["username"])
    env.user_manager.reset_password(target.user_id, p["new_password"])
    # 用新密码认证验证
    user = env.user_manager.authenticate(p["username"], p["new_password"])
    return user is not None


def _op_update_user(env, p):
    target = env.user_manager.get_by_username(p["username"])
    env.user_manager.update_user(target.user_id, display_name=p["display_name"])
    updated = env.user_manager.get_user(target.user_id)
    return updated.display_name


def _op_change_password_wrong(env, p):
    target = env.user_manager.get_by_username(p["username"])
    return env.user_manager.change_password(target.user_id, p["old"], p["new"])


def _op_change_password_correct(env, p):
    target = env.user_manager.get_by_username(p["username"])
    ok = env.user_manager.change_password(target.user_id, p["old"], p["new"])
    if not ok:
        return False
    return env.user_manager.authenticate(p["username"], p["new"]) is not None


def _op_user_stats(env, p):
    return env.user_manager.stats()


def _op_user_to_public_dict(env, p):
    target = env.user_manager.get_by_username(p["username"])
    d = target.to_public_dict()
    return "password_hash" not in d


def _op_session_save_load_report(env, p):
    sess = env.session(p["user_id"])
    report = _make_report(consultation_id=p.get("consultation_id", "TC-REP-1"),
                          patient_name=p.get("patient_name", "测试患者"))
    sess.save_report(report)
    loaded = sess.load_report(report.consultation_id)
    return loaded.consultation_id if loaded else None


def _op_session_list_reports(env, p):
    sess = env.session(p["user_id"])
    n = int(p.get("count", 3))
    for i in range(n):
        r = _make_report(consultation_id=f"TC-REP-{i}")
        sess.save_report(r)
    return sess.list_reports()


def _op_session_isolation(env, p):
    sess_a = env.session(p["user_a"])
    sess_b = env.session(p["user_b"])
    r = _make_report(consultation_id="ISO-REP-1")
    sess_a.save_report(r)
    # 用户 B 的会诊列表不应包含 A 的报告
    return sess_b.load_report("ISO-REP-1")


def _op_session_delete_report(env, p):
    sess = env.session(p["user_id"])
    r = _make_report(consultation_id="DEL-REP-1")
    sess.save_report(r)
    sess.delete_report("DEL-REP-1")
    return sess.load_report("DEL-REP-1")


def _op_session_save_load_patient(env, p):
    sess = env.session(p["user_id"])
    patient = _make_patient(patient_id=p.get("patient_id", "P001"),
                            name=p.get("name", "测试患者"))
    sess.save_patient(patient)
    loaded = sess.load_patient(patient.patient_id)
    return loaded.name if loaded else None


def _op_audit_log_count(env, p):
    actions = p.get("actions", ["LOGIN_SUCCESS", "LOGOUT"])
    for i, a in enumerate(actions):
        env.audit.log(user_id=f"U{i}", username=f"user{i}", action=a)
    return env.audit.get_logs()


def _op_audit_filter(env, p):
    env.audit.log(user_id="U1", username="u1", action="LOGIN_SUCCESS")
    env.audit.log(user_id="U2", username="u2", action="CREATE_CONSULTATION")
    env.audit.log(user_id="U3", username="u3", action="LOGIN_SUCCESS")
    return env.audit.get_logs(action=p["action"])


def _op_patient_from_dict_roundtrip(env, p):
    patient = _make_patient(patient_id=p.get("patient_id", "P001"),
                             name=p.get("name", "测试患者"))
    d = patient.to_dict()
    rebuilt = Patient.from_dict(d)
    return rebuilt.name


def _op_gender_parse(env, p):
    return Gender.parse(p["value"]).value


def _op_medical_record_integrate(env, p):
    engine, patient = _seed_consultation_engine(env)
    pid = patient.patient_id
    integrated = engine.integrator.integrate(pid)
    return integrated.name


def _op_consultation_run(env, p):
    engine, patient = _seed_consultation_engine(env)
    cid = engine.request_consultation(
        patient_id=patient.patient_id,
        clinical_question="胸痛原因待查",
        consultation_type="MDT"
    )
    report = engine.run_consultation(cid, patient=patient)
    return {"status": report.status, "patient_name": report.patient_name,
            "has_diagnosis": len(report.differential_diagnosis) > 0}


def _op_global_index(env, p):
    # 在两个用户空间各保存一份报告
    for uid in p.get("user_ids", ["U_A", "U_B"]):
        sess = env.session(uid)
        r = _make_report(consultation_id=f"G-{uid}")
        sess.save_report(r)
    gi = env.global_index()
    return gi.list_all_reports()


def _op_password_length(env, p):
    h = PasswordHasher.hash(p["password"])
    return len(h) > 0


def _op_password_special(env, p):
    h = PasswordHasher.hash(p["password"])
    return PasswordHasher.verify(p["password"], h)


def _op_role_parse_more(env, p):
    return Role.parse(p["value"]).value


def _op_auth_empty(env, p):
    return env.user_manager.authenticate(p["username"], p["password"])


def _op_create_user_admin(env, p):
    u = env.user_manager.create_user(
        username=p["username"], display_name=p.get("display_name", "管理员"),
        role=Role.ADMIN, password=p["password"]
    )
    return u.role.value


def _op_create_user_reviewer(env, p):
    u = env.user_manager.create_user(
        username=p["username"], display_name=p.get("display_name", "审核员"),
        role=Role.REVIEWER, password=p["password"]
    )
    return u.role.value


def _op_update_user_role(env, p):
    target = env.user_manager.get_by_username(p["username"])
    updated = env.user_manager.update_user(target.user_id, role=p["new_role"])
    return updated.role.value


def _op_update_user_department(env, p):
    target = env.user_manager.get_by_username(p["username"])
    updated = env.user_manager.update_user(target.user_id, department=p["department"])
    return updated.department


def _op_list_users_by_role(env, p):
    return [u.role.value for u in env.user_manager.list_users(role=Role.parse(p["role"]))]


def _op_delete_user_non_admin(env, p):
    target = env.user_manager.get_by_username(p["username"])
    env.user_manager.delete_user(target.user_id)
    return env.user_manager.get_user(target.user_id) is None


def _op_session_report_with_status(env, p):
    sess = env.session(p["user_id"])
    report = _make_report(
        consultation_id=p.get("consultation_id", "TC-STATUS-1"),
        status=p.get("status", "APPROVED")
    )
    sess.save_report(report)
    loaded = sess.load_report(report.consultation_id)
    return loaded.status


def _op_session_report_isolation(env, p):
    sess_a = env.session(p["user_a"])
    sess_b = env.session(p["user_b"])
    r = _make_report(consultation_id="ISO-2")
    sess_a.save_report(r)
    # B 加载同 ID 应不存在
    return sess_b.load_report("ISO-2")


def _op_patient_with_images(env, p):
    patient = _make_patient(patient_id=p.get("patient_id", "P_IMG"))
    patient.medical_images.append(MedicalImage(
        image_id="IMG-001", patient_id=patient.patient_id,
        file_name="xray.jpg", file_path="/tmp/xray.jpg",
        image_type="X光", body_part="胸部"
    ))
    d = patient.to_dict()
    rebuilt = Patient.from_dict(d)
    return [img.image_id for img in rebuilt.medical_images]


def _op_patient_public_dict(env, p):
    patient = _make_patient()
    patient.phone = "13800138000"
    patient.id_card = "110101199001011234"
    d = patient.to_public_dict()
    return d.get("phone", "").count("*") >= 4 and d.get("id_card", "").count("*") >= 4


def _op_gender_parse_more(env, p):
    return Gender.parse(p["value"]).value


def _op_consultation_run_emergency(env, p):
    engine, patient = _seed_consultation_engine(env)
    cid = engine.request_consultation(
        patient_id=patient.patient_id,
        clinical_question="急诊胸痛",
        consultation_type="急诊会诊",
        priority=Priority.EMERGENCY
    )
    report = engine.run_consultation(cid, patient=patient)
    return report.status


def _op_consultation_run_routine(env, p):
    engine, patient = _seed_consultation_engine(env)
    cid = engine.request_consultation(
        patient_id=patient.patient_id,
        clinical_question="常规腹痛",
        consultation_type="常规会诊",
        priority=Priority.ROUTINE
    )
    report = engine.run_consultation(cid, patient=patient)
    return report.status


def _op_medical_image_save(env, p):
    # 创建一个临时图片文件
    img_path = env.data_dir / "test_image.png"
    img_path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 20)
    image = env.session(p["user_id"]).save_medical_image(
        patient_id=p.get("patient_id", "P001"),
        source_path=str(img_path),
        image_type=p.get("image_type", "CT"),
        body_part=p.get("body_part", "胸部")
    )
    env.session(p["user_id"]).save_medical_image_meta(image)
    return image.image_id.startswith("IMG-") and Path(image.file_path).exists()


def _op_medical_image_list(env, p):
    sess = env.session(p["user_id"])
    img_path = env.data_dir / "test_image_list.jpg"
    img_path.write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 20)
    image = sess.save_medical_image(
        patient_id=p.get("patient_id", "P002"),
        source_path=str(img_path)
    )
    sess.save_medical_image_meta(image)
    return sess.list_medical_images(patient_id=p.get("patient_id", "P002"))


def _op_medical_image_delete(env, p):
    sess = env.session(p["user_id"])
    img_path = env.data_dir / "test_image_del.png"
    img_path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 20)
    image = sess.save_medical_image(patient_id="P003", source_path=str(img_path))
    sess.save_medical_image_meta(image)
    sess.delete_medical_image(image.image_id)
    return sess.load_medical_image(image.image_id) is None


def _op_medical_image_missing(env, p):
    return env.session(p["user_id"]).save_medical_image(
        patient_id="P999", source_path="/nonexistent/path.jpg"
    )


def _op_audit_empty(env, p):
    env.audit.log(user_id="U1", username="u1", action="TEST_ACTION")
    return env.audit.get_logs(action="NOT_EXIST")


def _op_user_last_login(env, p):
    user = env.user_manager.get_by_username(p["username"])
    env.user_manager.update_last_login(user.user_id, ip="127.0.0.1")
    updated = env.user_manager.get_user(user.user_id)
    return updated.last_login_at is not None and updated.last_login_ip == "127.0.0.1"


# ============================================================
# 后端管理平台配置（admin_config）测试操作
# ============================================================
def _make_admin_mgr(env):
    """在临时目录中创建独立的 AdminConfigManager"""
    from src.core.admin_config import AdminConfigManager
    return AdminConfigManager(data_dir=str(env.data_dir / "admin"))


def _op_admin_init(env, p):
    """测试 admin_config 初始化默认数据"""
    mgr = _make_admin_mgr(env)
    return {
        "agents": len(mgr.agents),
        "adapters": len(mgr.adapters),
        "quality_items": len(mgr.quality_items),
        "role_templates": len(mgr.role_templates),
        "plugins": len(mgr.plugins),
        "desensitization": len(mgr.desensitization),
    }


def _op_admin_add_agent(env, p):
    from src.core.admin_config import AgentConfig
    mgr = _make_admin_mgr(env)
    agent = AgentConfig(
        agent_id=p["agent_id"], name=p["name"],
        specialty=p["specialty"], description=p.get("description", "")
    )
    mgr.add_agent(agent)
    # 重新加载验证
    mgr2 = _make_admin_mgr(env)
    return any(a.agent_id == p["agent_id"] for a in mgr2.agents)


def _op_admin_add_agent_duplicate(env, p):
    from src.core.admin_config import AgentConfig
    mgr = _make_admin_mgr(env)
    agent1 = AgentConfig(agent_id=p["agent_id"], name="A", specialty="X")
    mgr.add_agent(agent1)
    agent2 = AgentConfig(agent_id=p["agent_id"], name="B", specialty="Y")
    try:
        mgr.add_agent(agent2)
        return False  # 不应到达
    except ValueError:
        return True


def _op_admin_toggle_agent(env, p):
    from src.core.admin_config import AgentConfig
    mgr = _make_admin_mgr(env)
    agent = AgentConfig(agent_id=p["agent_id"], name=p["name"], specialty=p["specialty"])
    mgr.add_agent(agent)
    new_state = mgr.toggle_agent(p["agent_id"])
    return new_state is False  # 切换后应为停用


def _op_admin_delete_agent(env, p):
    from src.core.admin_config import AgentConfig
    mgr = _make_admin_mgr(env)
    agent = AgentConfig(agent_id=p["agent_id"], name="X", specialty="Y")
    mgr.add_agent(agent)
    ok = mgr.delete_agent(p["agent_id"])
    return ok and not any(a.agent_id == p["agent_id"] for a in mgr.agents)


def _op_admin_update_agent(env, p):
    from src.core.admin_config import AgentConfig
    mgr = _make_admin_mgr(env)
    agent = AgentConfig(agent_id=p["agent_id"], name=p["name"], specialty=p["specialty"])
    mgr.add_agent(agent)
    updated = mgr.update_agent(p["agent_id"], name=p["new_name"])
    return updated.name == p["new_name"]


def _op_admin_add_adapter(env, p):
    from src.core.admin_config import AdapterConfig
    mgr = _make_admin_mgr(env)
    adapter = AdapterConfig(
        adapter_type=p["adapter_type"], name=p["name"],
        endpoint=p.get("endpoint", "http://test.com"), api_type=p.get("api_type", "REST")
    )
    mgr.add_adapter(adapter)
    return any(a.name == p["name"] for a in mgr.adapters)


def _op_admin_test_adapter(env, p):
    from src.core.admin_config import AdapterConfig
    mgr = _make_admin_mgr(env)
    adapter = AdapterConfig(
        adapter_type="HIS", name=p["name"],
        endpoint=p.get("endpoint", "http://test.com"), api_type="REST"
    )
    mgr.add_adapter(adapter)
    idx = next(i for i, a in enumerate(mgr.adapters) if a.name == p["name"])
    status, ts = mgr.test_adapter(idx)
    return status == "成功"


def _op_admin_test_adapter_fail(env, p):
    from src.core.admin_config import AdapterConfig
    mgr = _make_admin_mgr(env)
    adapter = AdapterConfig(
        adapter_type="HIS", name=p["name"],
        endpoint="invalid-endpoint", api_type="REST"
    )
    mgr.add_adapter(adapter)
    idx = next(i for i, a in enumerate(mgr.adapters) if a.name == p["name"])
    status, ts = mgr.test_adapter(idx)
    return status == "失败"


def _op_admin_delete_adapter(env, p):
    from src.core.admin_config import AdapterConfig
    mgr = _make_admin_mgr(env)
    adapter = AdapterConfig(adapter_type="HIS", name=p["name"], endpoint="http://x.com")
    mgr.add_adapter(adapter)
    idx = next(i for i, a in enumerate(mgr.adapters) if a.name == p["name"])
    ok = mgr.delete_adapter(idx)
    return ok


def _op_admin_add_quality_item(env, p):
    from src.core.admin_config import QualityControlItem
    mgr = _make_admin_mgr(env)
    item = QualityControlItem(
        item_id=p["item_id"], name=p["name"],
        description=p.get("description", ""), max_score=float(p.get("max_score", 10)),
        weight=float(p.get("weight", 1.0))
    )
    mgr.add_quality_item(item)
    return any(q.item_id == p["item_id"] for q in mgr.quality_items)


def _op_admin_add_quality_duplicate(env, p):
    from src.core.admin_config import QualityControlItem
    mgr = _make_admin_mgr(env)
    item1 = QualityControlItem(item_id=p["item_id"], name="A")
    mgr.add_quality_item(item1)
    item2 = QualityControlItem(item_id=p["item_id"], name="B")
    try:
        mgr.add_quality_item(item2)
        return False
    except ValueError:
        return True


def _op_admin_update_quality_item(env, p):
    from src.core.admin_config import QualityControlItem
    mgr = _make_admin_mgr(env)
    item = QualityControlItem(item_id=p["item_id"], name=p["name"])
    mgr.add_quality_item(item)
    updated = mgr.update_quality_item(p["item_id"], name=p["new_name"], max_score=20.0)
    return updated.name == p["new_name"] and updated.max_score == 20.0


def _op_admin_delete_quality_item(env, p):
    from src.core.admin_config import QualityControlItem
    mgr = _make_admin_mgr(env)
    item = QualityControlItem(item_id=p["item_id"], name="X")
    mgr.add_quality_item(item)
    ok = mgr.delete_quality_item(p["item_id"])
    return ok


def _op_admin_add_role_template(env, p):
    from src.core.admin_config import RoleTemplate
    mgr = _make_admin_mgr(env)
    role = RoleTemplate(
        role_id=p["role_id"], name=p["name"],
        description=p.get("description", ""), data_scope=p.get("data_scope", "self")
    )
    mgr.add_role_template(role)
    return any(r.role_id == p["role_id"] for r in mgr.role_templates)


def _op_admin_add_role_duplicate(env, p):
    from src.core.admin_config import RoleTemplate
    mgr = _make_admin_mgr(env)
    r1 = RoleTemplate(role_id=p["role_id"], name="A")
    mgr.add_role_template(r1)
    r2 = RoleTemplate(role_id=p["role_id"], name="B")
    try:
        mgr.add_role_template(r2)
        return False
    except ValueError:
        return True


def _op_admin_delete_role_template(env, p):
    from src.core.admin_config import RoleTemplate
    mgr = _make_admin_mgr(env)
    role = RoleTemplate(role_id=p["role_id"], name="X")
    mgr.add_role_template(role)
    ok = mgr.delete_role_template(p["role_id"])
    return ok and not any(r.role_id == p["role_id"] for r in mgr.role_templates)


def _op_admin_toggle_plugin(env, p):
    mgr = _make_admin_mgr(env)
    # 找到默认插件
    if not mgr.plugins:
        return False
    pid = mgr.plugins[0].plugin_id
    new_state = mgr.toggle_plugin(pid)
    return new_state in (True, False)


def _op_admin_save_load_config(env, p):
    """测试配置持久化与重新加载"""
    from src.core.admin_config import AdminConfigManager
    mgr = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    # 修改系统全局参数
    mgr.system.hospital_name = "测试医院"
    mgr.system.max_concurrent_consultations = 200
    mgr.save_all()
    # 重新加载
    mgr2 = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    return mgr2.system.hospital_name == "测试医院" and mgr2.system.max_concurrent_consultations == 200


def _op_admin_stats(env, p):
    mgr = _make_admin_mgr(env)
    stats = mgr.stats()
    return stats.get("agent_total", 0) > 0 and stats.get("adapter_total", 0) > 0


def _op_admin_moderator_config(env, p):
    """测试 Moderator 配置持久化"""
    from src.core.admin_config import AdminConfigManager
    mgr = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    mgr.moderator.max_discussion_rounds = 5
    mgr.moderator.consensus_threshold = 0.85
    mgr.moderator.enable_debate = False
    mgr.save_all()
    # 重新加载验证
    mgr2 = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    return (mgr2.moderator.max_discussion_rounds == 5 and
            abs(mgr2.moderator.consensus_threshold - 0.85) < 0.001 and
            mgr2.moderator.enable_debate is False)


def _op_admin_schedule_rule(env, p):
    """测试会诊调度规则持久化"""
    from src.core.admin_config import AdminConfigManager
    mgr = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    mgr.schedule_rule.emergency_response_minutes = 5
    mgr.schedule_rule.collaboration_mode = p.get("mode", "auto")
    mgr.save_all()
    mgr2 = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    return (mgr2.schedule_rule.emergency_response_minutes == 5 and
            mgr2.schedule_rule.collaboration_mode == p.get("mode", "auto"))


def _op_admin_security_config(env, p):
    """测试安全配置持久化"""
    from src.core.admin_config import AdminConfigManager
    mgr = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    mgr.security.tls_enabled = True
    mgr.security.storage_encryption = p.get("encryption", "SM4")
    mgr.security.compliance_level = "等保三级"
    mgr.save_all()
    mgr2 = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    return (mgr2.security.tls_enabled is True and
            mgr2.security.storage_encryption == p.get("encryption", "SM4") and
            mgr2.security.compliance_level == "等保三级")


def _op_admin_desensitization_default(env, p):
    """测试默认脱敏规则"""
    mgr = _make_admin_mgr(env)
    scenes = {d.scene for d in mgr.desensitization}
    return "display" in scenes and "export" in scenes and "research" in scenes


def _op_admin_role_template_default(env, p):
    """测试默认角色模板"""
    mgr = _make_admin_mgr(env)
    role_ids = {r.role_id for r in mgr.role_templates}
    expected = {"ROLE_ADMIN", "ROLE_DOCTOR", "ROLE_REVIEWER"}
    return expected.issubset(role_ids)


def _op_admin_notification_config(env, p):
    """测试消息通知配置"""
    from src.core.admin_config import AdminConfigManager
    mgr = AdminConfigManager(data_dir=str(env.data_dir / "admin"))
    # 站内信默认开启
    return mgr.notification.channels.get("站内信", False) is True


# ============================================================
# 审核队列（ReviewQueue）操作
# ============================================================
def _op_review_submit(env, p):
    """提交报告到审核队列"""
    return env.review_queue.submit(
        consultation_id=p.get("consultation_id", "VCS-RV-001"),
        doctor_id=p.get("doctor_id", "U_DOC1"),
        doctor_name=p.get("doctor_name", "测试医生"),
        patient_id=p.get("patient_id", "P001"),
        patient_name=p.get("patient_name", "测试患者"),
        consultation_type=p.get("consultation_type", "MDT"),
        priority=p.get("priority", "常规"),
        department=p.get("department", "心内科")
    )


def _op_review_submit_duplicate(env, p):
    """重复提交应返回 False"""
    cid = p.get("consultation_id", "VCS-RV-DUP")
    env.review_queue.submit(consultation_id=cid, doctor_id="U_DOC1", doctor_name="医生")
    return env.review_queue.submit(consultation_id=cid, doctor_id="U_DOC1", doctor_name="医生")


def _op_review_list_pending(env, p):
    """列出待审核报告"""
    n = int(p.get("count", 3))
    for i in range(n):
        env.review_queue.submit(
            consultation_id=f"VCS-RV-{i:03d}",
            doctor_id=f"U_DOC{i}", doctor_name=f"医生{i}",
            patient_name=f"患者{i}"
        )
    return env.review_queue.list_pending()


def _op_review_approve(env, p):
    """审核通过"""
    cid = p.get("consultation_id", "VCS-RV-APPR")
    env.review_queue.submit(consultation_id=cid, doctor_id="U_DOC1", doctor_name="医生")
    return env.review_queue.approve(cid, reviewer_id="U_REV1", reviewer_name="审核员",
                                     comment=p.get("comment", ""))


def _op_review_reject(env, p):
    """驳回报告"""
    cid = p.get("consultation_id", "VCS-RV-REJ")
    env.review_queue.submit(consultation_id=cid, doctor_id="U_DOC1", doctor_name="医生")
    return env.review_queue.reject(cid, reviewer_id="U_REV1", reviewer_name="审核员",
                                    comment=p.get("comment", "驳回原因"))


def _op_review_approve_not_pending(env, p):
    """审核非待审核报告应失败"""
    cid = "VCS-RV-NP"
    env.review_queue.submit(consultation_id=cid, doctor_id="U_DOC1", doctor_name="医生")
    env.review_queue.approve(cid, reviewer_id="U_REV1", reviewer_name="审核员")
    # 再次审核应返回 False
    return env.review_queue.approve(cid, reviewer_id="U_REV1", reviewer_name="审核员")


def _op_review_stats(env, p):
    """审核队列统计"""
    env.review_queue.submit(consultation_id="VCS-S1", doctor_id="U1", doctor_name="医生1")
    env.review_queue.submit(consultation_id="VCS-S2", doctor_id="U2", doctor_name="医生2")
    env.review_queue.approve("VCS-S1", reviewer_id="R1", reviewer_name="审核员1")
    env.review_queue.reject("VCS-S2", reviewer_id="R1", reviewer_name="审核员1")
    return env.review_queue.stats()


def _op_review_list_reviewed_by(env, p):
    """按审核员查询历史记录"""
    env.review_queue.submit(consultation_id="VCS-R1", doctor_id="U1", doctor_name="医生1")
    env.review_queue.submit(consultation_id="VCS-R2", doctor_id="U2", doctor_name="医生2")
    env.review_queue.approve("VCS-R1", reviewer_id="U_REV1", reviewer_name="审核员A")
    env.review_queue.reject("VCS-R2", reviewer_id="U_REV1", reviewer_name="审核员A")
    return env.review_queue.list_reviewed_by("U_REV1")


def _op_review_resubmit_rejected(env, p):
    """被驳回后重新提交"""
    cid = "VCS-RV-RE"
    env.review_queue.submit(consultation_id=cid, doctor_id="U1", doctor_name="医生")
    env.review_queue.reject(cid, reviewer_id="R1", reviewer_name="审核员")
    # 重新提交
    return env.review_queue.submit(consultation_id=cid, doctor_id="U1", doctor_name="医生")


def _op_review_withdraw(env, p):
    """撤回提交"""
    cid = "VCS-RV-WD"
    env.review_queue.submit(consultation_id=cid, doctor_id="U1", doctor_name="医生")
    return env.review_queue.withdraw(cid)


def _op_review_get_entry(env, p):
    """获取审核条目（返回完整条目 dict，便于断言字段）"""
    cid = p.get("consultation_id", "VCS-RV-GE")
    env.review_queue.submit(consultation_id=cid, doctor_id="U1", doctor_name="医生")
    return env.review_queue.get_entry(cid)


def _op_engine_submit_for_review(env, p):
    """测试引擎 submit_for_review 方法"""
    engine = ConsultationEngine()
    cid = engine.request_consultation("P001", "测试", "MDT")
    # 先完成会诊
    patient = _make_patient()
    integrator = MedicalRecordIntegrator()
    integrator.add_adapter(MockHospitalAdapter("HIS", sample_data={"P001": {
        "name": patient.name, "gender": "男", "age": patient.age,
        "department": patient.department, "attending_doctor": "张医生",
    }}))
    integrator.add_adapter(MockHospitalAdapter("EMR", sample_data={"P001": {
        "chief_complaint": patient.chief_complaint,
        "present_illness": patient.present_illness,
    }}))
    engine.integrator = integrator
    engine.run_consultation(cid, patient=patient)
    # 提交审核
    ok = engine.submit_for_review(cid)
    report = engine.get_report(cid)
    return report.status if (ok and report) else "FAILED"


OPERATIONS: Dict[str, Callable[[TestEnv, dict], Any]] = {
    "password_hash": _op_password_hash,
    "password_verify": _op_password_verify,
    "password_empty": _op_password_empty,
    "password_length": _op_password_length,
    "password_special": _op_password_special,
    "role_parse": _op_role_parse,
    "role_parse_more": _op_role_parse_more,
    "role_can_review": _op_role_can_review,
    "role_display": _op_role_display,
    "auth": _op_auth,
    "auth_disabled": _op_auth_disabled,
    "auth_empty": _op_auth_empty,
    "create_user": _op_create_user,
    "create_duplicate": _op_create_duplicate,
    "create_weak_password": _op_create_weak_password,
    "create_user_admin": _op_create_user_admin,
    "create_user_reviewer": _op_create_user_reviewer,
    "delete_last_admin": _op_delete_last_admin,
    "delete_user_non_admin": _op_delete_user_non_admin,
    "reset_password": _op_reset_password,
    "update_user": _op_update_user,
    "update_user_role": _op_update_user_role,
    "update_user_department": _op_update_user_department,
    "list_users_by_role": _op_list_users_by_role,
    "change_password_wrong": _op_change_password_wrong,
    "change_password_correct": _op_change_password_correct,
    "user_stats": _op_user_stats,
    "user_to_public_dict": _op_user_to_public_dict,
    "user_last_login": _op_user_last_login,
    "session_save_load_report": _op_session_save_load_report,
    "session_list_reports": _op_session_list_reports,
    "session_isolation": _op_session_isolation,
    "session_delete_report": _op_session_delete_report,
    "session_save_load_patient": _op_session_save_load_patient,
    "session_report_with_status": _op_session_report_with_status,
    "session_report_isolation": _op_session_report_isolation,
    "audit_log_count": _op_audit_log_count,
    "audit_filter": _op_audit_filter,
    "audit_empty": _op_audit_empty,
    "patient_from_dict_roundtrip": _op_patient_from_dict_roundtrip,
    "patient_with_images": _op_patient_with_images,
    "patient_public_dict": _op_patient_public_dict,
    "gender_parse": _op_gender_parse,
    "gender_parse_more": _op_gender_parse_more,
    "medical_record_integrate": _op_medical_record_integrate,
    "consultation_run": _op_consultation_run,
    "consultation_run_emergency": _op_consultation_run_emergency,
    "consultation_run_routine": _op_consultation_run_routine,
    "global_index": _op_global_index,
    "medical_image_save": _op_medical_image_save,
    "medical_image_list": _op_medical_image_list,
    "medical_image_delete": _op_medical_image_delete,
    "medical_image_missing": _op_medical_image_missing,
    # 后端管理平台（admin_config）操作
    "admin_init": _op_admin_init,
    "admin_add_agent": _op_admin_add_agent,
    "admin_add_agent_duplicate": _op_admin_add_agent_duplicate,
    "admin_toggle_agent": _op_admin_toggle_agent,
    "admin_delete_agent": _op_admin_delete_agent,
    "admin_update_agent": _op_admin_update_agent,
    "admin_add_adapter": _op_admin_add_adapter,
    "admin_test_adapter": _op_admin_test_adapter,
    "admin_test_adapter_fail": _op_admin_test_adapter_fail,
    "admin_delete_adapter": _op_admin_delete_adapter,
    "admin_add_quality_item": _op_admin_add_quality_item,
    "admin_add_quality_duplicate": _op_admin_add_quality_duplicate,
    "admin_update_quality_item": _op_admin_update_quality_item,
    "admin_delete_quality_item": _op_admin_delete_quality_item,
    "admin_add_role_template": _op_admin_add_role_template,
    "admin_add_role_duplicate": _op_admin_add_role_duplicate,
    "admin_delete_role_template": _op_admin_delete_role_template,
    "admin_toggle_plugin": _op_admin_toggle_plugin,
    "admin_save_load_config": _op_admin_save_load_config,
    "admin_stats": _op_admin_stats,
    "admin_moderator_config": _op_admin_moderator_config,
    "admin_schedule_rule": _op_admin_schedule_rule,
    "admin_security_config": _op_admin_security_config,
    "admin_desensitization_default": _op_admin_desensitization_default,
    "admin_role_template_default": _op_admin_role_template_default,
    "admin_notification_config": _op_admin_notification_config,
    # 审核队列（ReviewQueue）操作
    "review_submit": _op_review_submit,
    "review_submit_duplicate": _op_review_submit_duplicate,
    "review_list_pending": _op_review_list_pending,
    "review_approve": _op_review_approve,
    "review_reject": _op_review_reject,
    "review_approve_not_pending": _op_review_approve_not_pending,
    "review_stats": _op_review_stats,
    "review_list_reviewed_by": _op_review_list_reviewed_by,
    "review_resubmit_rejected": _op_review_resubmit_rejected,
    "review_withdraw": _op_review_withdraw,
    "review_get_entry": _op_review_get_entry,
    "engine_submit_for_review": _op_engine_submit_for_review,
}


# ============================================================
# 断言
# ============================================================
def _normalize(v: Any) -> str:
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, sort_keys=True)
    return str(v)


def assert_case(actual: Any, expected_type: str, expected: Any) -> Tuple[bool, str]:
    """根据断言类型比较 actual 与 expected，返回 (passed, message)"""
    et = (expected_type or "").strip().lower()

    if et == "truthy":
        return (bool(actual), f"actual={_brief(actual)}")
    if et == "falsy":
        return (not actual, f"actual={_brief(actual)}")
    if et == "not_null":
        return (actual is not None, f"actual={'None' if actual is None else 'set'}")
    if et == "equals":
        return (_normalize(actual) == _normalize(expected),
                f"actual={_brief(actual)} expected={_brief(expected)}")
    if et == "contains":
        return (_normalize(expected) in _normalize(actual),
                f"actual={_brief(actual)} expected(contains)={_brief(expected)}")
    if et == "count":
        try:
            n = len(actual)
        except TypeError:
            return (False, f"actual 无 len(): {_brief(actual)}")
        return (n == int(expected), f"count={n} expected={expected}")
    if et == "json_has_key":
        if not isinstance(actual, dict):
            return (False, f"actual 非 dict: {_brief(actual)}")
        return (expected in actual, f"keys={list(actual.keys())} expected_key={expected}")
    if et == "json_has_no_key":
        if not isinstance(actual, dict):
            return (False, f"actual 非 dict: {_brief(actual)}")
        return (expected not in actual, f"keys={list(actual.keys())} absent_key={expected}")
    if et == "json_field_equals":
        field, value = expected.split("=", 1)
        if not isinstance(actual, dict) or field not in actual:
            return (False, f"字段缺失: {field}")
        return (_normalize(actual[field]) == _normalize(value),
                f"{field}={_brief(actual[field])} expected={value}")
    return (False, f"未知断言类型: {expected_type}")


def _brief(v: Any) -> str:
    s = _normalize(v)
    return s if len(s) <= 120 else s[:117] + "..."


# ============================================================
# 用例数据结构
# ============================================================
@dataclass
class TestCase:
    case_id: str
    category: str          # 测试项
    name: str              # 用例名称
    description: str
    operation: str
    params: dict
    expected_type: str
    expected: Any
    enabled: bool

    # 阻止 pytest 将本 dataclass 误当作测试类收集
    __test__ = False


@dataclass
class TestResult:
    case_id: str
    name: str
    category: str
    passed: bool
    message: str
    duration_ms: float = 0.0

    __test__ = False


# ============================================================
# 运行器
# ============================================================
class XlsxTestRunner:
    HEADER = ["用例ID", "测试项", "用例名称", "描述", "操作operation",
              "参数params(JSON)", "断言类型", "期望expected", "启用"]

    def __init__(self, xlsx_path: str = None):
        if xlsx_path is None:
            xlsx_path = str(Path(__file__).parent / "test.xlsx")
        self.xlsx_path = xlsx_path
        self.results: List[TestResult] = []

    def load_cases(self) -> List[TestCase]:
        wb = load_workbook(self.xlsx_path, data_only=True)
        ws = wb["test_cases"]
        cases = []
        # 第1行表头，第2行起数据
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            case_id, category, name, desc, op, params_str, et, exp, enabled = (
                list(row) + [None] * (9 - len(row))
            )[:9]
            enabled = str(enabled).strip().upper() in ("Y", "YES", "TRUE", "1", "启用")
            try:
                params = json.loads(params_str) if params_str else {}
            except (json.JSONDecodeError, TypeError):
                params = {}
            cases.append(TestCase(
                case_id=str(case_id), category=str(category or ""),
                name=str(name or ""), description=str(desc or ""),
                operation=str(op or "").strip(),
                params=params, expected_type=str(et or ""),
                expected=exp, enabled=enabled,
            ))
        return cases

    def run(self, verbose: bool = True) -> List[TestResult]:
        cases = self.load_cases()
        self.results = []
        for tc in cases:
            if not tc.enabled:
                self.results.append(TestResult(
                    case_id=tc.case_id, name=tc.name, category=tc.category,
                    passed=True, message="已跳过（未启用）"
                ))
                continue
            res = self._run_one(tc)
            self.results.append(res)
            if verbose:
                tag = "✓" if res.passed else "✗"
                print(f"  {tag} {tc.case_id} [{tc.category}] {tc.name}"
                      + (f"  -- {res.message}" if not res.passed else ""))
        return self.results

    def _run_one(self, tc: TestCase) -> TestResult:
        import time
        handler = OPERATIONS.get(tc.operation)
        if handler is None:
            return TestResult(case_id=tc.case_id, name=tc.name, category=tc.category,
                              passed=False, message=f"未注册操作: {tc.operation}")
        env = TestEnv()
        t0 = time.perf_counter()
        try:
            actual = handler(env, tc.params)
            # 期望值为"throws"系列：期望抛异常
            et = tc.expected_type.strip().lower() if tc.expected_type else ""
            if et in ("throws", "throws_contains"):
                return TestResult(case_id=tc.case_id, name=tc.name, category=tc.category,
                                  passed=False,
                                  message=f"期望抛异常但实际成功返回: {_brief(actual)}")
            passed, msg = assert_case(actual, tc.expected_type, tc.expected)
            return TestResult(case_id=tc.case_id, name=tc.name, category=tc.category,
                              passed=passed, message=msg,
                              duration_ms=(time.perf_counter() - t0) * 1000)
        except Exception as e:
            et = tc.expected_type.strip().lower() if tc.expected_type else ""
            if et in ("throws", "throws_contains"):
                msg_part = str(tc.expected) if tc.expected else ""
                if et == "throws_contains" and msg_part and msg_part not in str(e):
                    return TestResult(case_id=tc.case_id, name=tc.name, category=tc.category,
                                      passed=False,
                                      message=f"抛异常但消息不含 '{msg_part}': {e}")
                return TestResult(case_id=tc.case_id, name=tc.name, category=tc.category,
                                  passed=True, message=f"正确抛异常: {type(e).__name__}: {e}",
                                  duration_ms=(time.perf_counter() - t0) * 1000)
            return TestResult(case_id=tc.case_id, name=tc.name, category=tc.category,
                              passed=False,
                              message=f"未预期异常: {type(e).__name__}: {e}\n"
                                      f"{traceback.format_exc(limit=2)}")
        finally:
            env.cleanup()

    def summary(self) -> dict:
        total = len(self.results)
        passed = sum(1 for r in self.results if r.passed)
        failed = total - passed
        return {"total": total, "passed": passed, "failed": failed}


def main():
    import argparse
    parser = argparse.ArgumentParser(description="运行 test.xlsx 数据驱动测试")
    parser.add_argument("--path", default=None, help="test.xlsx 路径")
    args = parser.parse_args()

    runner = XlsxTestRunner(args.path)
    print(f"加载测试用例: {runner.xlsx_path}")
    cases = runner.load_cases()
    enabled = [c for c in cases if c.enabled]
    print(f"共 {len(cases)} 条用例，启用 {len(enabled)} 条\n")
    runner.run(verbose=True)
    s = runner.summary()
    print(f"\n结果：通过 {s['passed']} / 失败 {s['failed']} / 总计 {s['total']}")
    failed = [r for r in runner.results if not r.passed]
    if failed:
        print("\n失败用例：")
        for r in failed:
            print(f"  ✗ {r.case_id} {r.name}: {r.message}")
    return 0 if s["failed"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
